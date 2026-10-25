from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt


BASE_DIR = Path(__file__).resolve().parent
OUTPUT = BASE_DIR.parent / "rapport_Groupe13.pptx"
SLIDES = [
    ("Signalement citoyen des problèmes urbains", "Projet 13 · Groupe 13 · Technologies Web"),
    ("Contexte et besoin", "Problèmes urbains visibles, mais difficiles à signaler et à suivre."),
    ("Objectifs du projet", "Centraliser les signalements et rendre leur traitement transparent."),
    ("Utilisateurs et rôles", "Citoyens · agents municipaux · administrateurs des services."),
    ("Périmètre fonctionnel", "Dépôt photo · suivi par code · liste filtrable · espace agent."),
    ("Parcours citoyen", "Décrire · localiser · joindre une photo · suivre l’intervention."),
    ("Parcours agent", "Connexion · demandes du service · intervention · changement de statut."),
    ("Architecture technique", "HTML5 · CSS responsive · JavaScript · Flask · SQLite."),
    ("Modèle de données", "Catégories · quartiers · signalements · interventions · agents."),
    ("Sécurité et validation", "CSRF · mots de passe hachés · requêtes paramétrées · contrôle des photos."),
    ("Interface et accessibilité", "À compléter avec les captures ordinateur et téléphone."),
    ("API et statistiques", "Par catégorie · par quartier · délai moyen · retards · évolution hebdomadaire."),
    ("Tests et données de démonstration", "5 catégories · 12 quartiers · 200 signalements · 250 interventions · 8 agents."),
    ("Résultats et limites", "À compléter après les essais fonctionnels et la validation des utilisateurs."),
    ("Conclusion et démonstration", "Ajouter le bilan du groupe et les étapes de la démonstration vidéo."),
]


def main():
    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)
    forest = RGBColor(18, 59, 53)
    coral = RGBColor(226, 123, 81)
    paper = RGBColor(247, 248, 243)
    for index, (title, body) in enumerate(SLIDES):
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        background = slide.background.fill
        background.solid()
        background.fore_color.rgb = forest if index == 0 else paper
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.75), Inches(11.6), Inches(1.2))
        title_frame = title_box.text_frame
        title_frame.text = title
        title_run = title_frame.paragraphs[0].runs[0]
        title_run.font.name = "Aptos Display"
        title_run.font.size = Pt(34 if index == 0 else 30)
        title_run.font.bold = True
        title_run.font.color.rgb = RGBColor(255, 255, 255) if index == 0 else forest
        body_box = slide.shapes.add_textbox(Inches(0.85), Inches(2.25), Inches(10.8), Inches(2.2))
        body_frame = body_box.text_frame
        body_frame.word_wrap = True
        body_frame.text = body
        body_run = body_frame.paragraphs[0].runs[0]
        body_run.font.name = "Aptos"
        body_run.font.size = Pt(23 if index == 0 else 21)
        body_run.font.color.rgb = RGBColor(215, 232, 109) if index == 0 else RGBColor(75, 91, 83)
        accent = slide.shapes.add_shape(1, Inches(0.85), Inches(6.8), Inches(1.1), Inches(0.08))
        accent.fill.solid()
        accent.fill.fore_color.rgb = coral
        accent.line.fill.background()
        footer = slide.shapes.add_textbox(Inches(10.7), Inches(6.75), Inches(1.8), Inches(0.3))
        footer.text_frame.text = f"GROUPE 13  ·  {index + 1:02d}"
        footer_run = footer.text_frame.paragraphs[0].runs[0]
        footer_run.font.size = Pt(10)
        footer_run.font.bold = True
        footer_run.font.color.rgb = RGBColor(255, 255, 255) if index == 0 else forest
    presentation.save(OUTPUT)
    print(f"Rapport de travail généré : {OUTPUT} ({len(SLIDES)} diapositives)")


if __name__ == "__main__":
    main()